import io
import os
import secrets
import psycopg
import time
import uuid
from datetime import datetime
from functools import wraps

import pyotp
import qrcode
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font
from flask import Flask, request, render_template, send_file, redirect, url_for, session
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)


app.secret_key = os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32))

QR_INTERVAL_SECONDS = 10
FORM_SESSION_TIMEOUT_SECONDS = 600

DEFAULT_QR_SIZE = 360
MIN_QR_INTERVAL_SECONDS = 3
MAX_QR_INTERVAL_SECONDS = 300
MIN_QR_SIZE = 120
MAX_QR_SIZE = 600
STATIC_QR_INTERVAL_SECONDS = 21600

DATABASE_URL = os.environ.get("DATABASE_URL")

if not DATABASE_URL:
    raise RuntimeError("DATABASE_URL environment variable is not set.")


if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)

def get_db():
    return psycopg.connect(DATABASE_URL)

active_sessions = {}


def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS class_sessions (
                id TEXT PRIMARY KEY,
                teacher_id INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                ended_at TEXT
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id SERIAL PRIMARY KEY,
                teacher_id INTEGER NOT NULL,
                class_session_id TEXT NOT NULL,
                full_name TEXT NOT NULL,
                student_id TEXT NOT NULL,
                timestamp TEXT NOT NULL,
                UNIQUE(teacher_id, student_id, class_session_id)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS pending_scans (
                session_id TEXT PRIMARY KEY,
                teacher_id INTEGER NOT NULL,
                class_session_id TEXT NOT NULL,
                created_at DOUBLE PRECISION NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS user_settings (
                teacher_id INTEGER PRIMARY KEY,
                qr_interval_seconds INTEGER NOT NULL DEFAULT 10,
                qr_size INTEGER NOT NULL DEFAULT 360,
                auto_refresh_enabled BOOLEAN NOT NULL DEFAULT TRUE
            )
        """)

        conn.execute("ALTER TABLE attendance ADD COLUMN IF NOT EXISTS crn TEXT")
        conn.execute("ALTER TABLE class_sessions ADD COLUMN IF NOT EXISTS deleted_at TEXT")

init_db()

def cleanup_expired_scans(conn):
    cutoff = time.time() - FORM_SESSION_TIMEOUT_SECONDS
    conn.execute("DELETE FROM pending_scans WHERE created_at < %s", (cutoff,))

def format_dt(iso_string):
    if not iso_string:
        return None
    return datetime.fromisoformat(iso_string).strftime("%d.%m.%Y %H:%M")

def get_user_by_id(user_id):
    conn = get_db()
    row = conn.execute(
        "SELECT id, name, email FROM users WHERE id = %s", (user_id,)
    ).fetchone()
    conn.close()
    if row is None:
        return None
    return {"id": row[0], "name": row[1], "email": row[2]}

def get_user_settings(user_id):
    conn = get_db()
    conn.execute(
        "INSERT INTO user_settings (teacher_id) VALUES (%s) ON CONFLICT (teacher_id) DO NOTHING",
        (user_id,),
    )
    conn.commit()
    row = conn.execute(
        "SELECT qr_interval_seconds, qr_size, auto_refresh_enabled FROM user_settings WHERE teacher_id = %s",
        (user_id,),
    ).fetchone()
    conn.close()
    return {
        "qr_interval_seconds": row[0],
        "qr_size": row[1],
        "auto_refresh_enabled": row[2],
    }


def current_user():
    user_id = session.get("user_id")
    if user_id is None:
        return None
    return get_user_by_id(user_id)

def login_required(view_func):
    @wraps(view_func)
    def wrapped(*args, **kwargs):
        if session.get("user_id") is None:
            return redirect(url_for("login"))
        return view_func(*args, **kwargs)
    return wrapped

@app.route("/")
def landing():
    return render_template("index.html", user=current_user())

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "GET":
        return render_template("register.html", error=None)

    name = request.form.get("name", "").strip()
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")

    if not name or not email or not password:
        return render_template("register.html", error="Lütfen tüm alanları doldurun.")

    if len(password) < 6:
        return render_template("register.html", error="Şifre en az 6 karakter olmalı.")

    conn = get_db()
    try:
        conn.execute(
            "INSERT INTO users (name, email, password_hash, created_at) VALUES (%s, %s, %s, %s)",
            (name, email, generate_password_hash(password), datetime.now().isoformat()),
        )
        conn.commit()
        user_id = conn.execute(
            "SELECT id FROM users WHERE email = %s", (email,)
        ).fetchone()[0]
    except psycopg.IntegrityError:
        conn.close()
        return render_template("register.html", error="Bu e-posta zaten kayıtlı.")
    conn.close()

    session["user_id"] = user_id
    return redirect(url_for("dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html", error=None)

    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")

    conn = get_db()
    row = conn.execute(
        "SELECT id, password_hash FROM users WHERE email = %s", (email,)
    ).fetchone()
    conn.close()

    if row is None or not check_password_hash(row[1], password):
        return render_template("login.html", error="E-posta veya şifre hatalı.")

    session["user_id"] = row[0]
    return redirect(url_for("dashboard"))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("landing"))

@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings_page():
    user = current_user()

    if request.method == "POST":
        try:
            interval = int(request.form.get("qr_interval_seconds", QR_INTERVAL_SECONDS))
        except ValueError:
            interval = QR_INTERVAL_SECONDS
        interval = max(MIN_QR_INTERVAL_SECONDS, min(MAX_QR_INTERVAL_SECONDS, interval))

        try:
            size = int(request.form.get("qr_size", DEFAULT_QR_SIZE))
        except ValueError:
            size = DEFAULT_QR_SIZE
        size = max(MIN_QR_SIZE, min(MAX_QR_SIZE, size))

        auto_refresh_enabled = request.form.get("auto_refresh_enabled") == "on"

        conn = get_db()
        conn.execute(
            "INSERT INTO user_settings (teacher_id, qr_interval_seconds, qr_size, auto_refresh_enabled) "
            "VALUES (%s, %s, %s, %s) "
            "ON CONFLICT (teacher_id) DO UPDATE SET "
            "qr_interval_seconds = EXCLUDED.qr_interval_seconds, "
            "qr_size = EXCLUDED.qr_size, "
            "auto_refresh_enabled = EXCLUDED.auto_refresh_enabled",
            (user["id"], interval, size, auto_refresh_enabled),
        )
        conn.commit()
        conn.close()

        return redirect(url_for("settings_page", saved=1))

    settings = get_user_settings(user["id"])
    return render_template(
        "settings.html",
        user=user,
        settings=settings,
        saved=request.args.get("saved") == "1",
        min_interval=MIN_QR_INTERVAL_SECONDS,
        max_interval=MAX_QR_INTERVAL_SECONDS,
        min_size=MIN_QR_SIZE,
        max_size=MAX_QR_SIZE,
    )


@app.route("/dashboard")
@login_required
def dashboard():
    user = current_user()
    active_info = active_sessions.get(user["id"])

    conn = get_db()
    session_rows = conn.execute(
        "SELECT id, created_at, ended_at FROM class_sessions "
        "WHERE teacher_id = %s AND deleted_at IS NULL ORDER BY created_at DESC",
        (user["id"],),
    ).fetchall()

    past_sessions = []
    for sid, created_at, ended_at in session_rows:
        count = conn.execute(
            "SELECT COUNT(*) FROM attendance WHERE teacher_id = %s AND class_session_id = %s",
            (user["id"], sid),
        ).fetchone()[0]
        is_active = active_info is not None and active_info["class_session_id"] == sid
        past_sessions.append({
            "id": sid,
            "created_at": format_dt(created_at),
            "ended_at": format_dt(ended_at),
            "count": count,
            "active": is_active,
        })

    conn.close()

    return render_template(
        "dashboard.html",
        user=user,
        has_active=active_info is not None,
        sessions=past_sessions,
        profile_saved=request.args.get("profile_saved") == "1",
    )

@app.route("/profile/edit", methods=["GET", "POST"])
@login_required
def edit_profile():
    user = current_user()

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")

        if not name or not email or not current_password:
            return render_template(
                "edit_profile.html", user=user,
                error="Please fill in your name, email, and current password.",
            )

        if new_password and len(new_password) < 6:
            return render_template(
                "edit_profile.html", user=user,
                error="New password must be at least 6 characters.",
            )

        conn = get_db()
        row = conn.execute(
            "SELECT password_hash FROM users WHERE id = %s", (user["id"],)
        ).fetchone()

        if row is None or not check_password_hash(row[0], current_password):
            conn.close()
            return render_template(
                "edit_profile.html", user=user,
                error="Current password is incorrect.",
            )

        try:
            if new_password:
                conn.execute(
                    "UPDATE users SET name = %s, email = %s, password_hash = %s WHERE id = %s",
                    (name, email, generate_password_hash(new_password), user["id"]),
                )
            else:
                conn.execute(
                    "UPDATE users SET name = %s, email = %s WHERE id = %s",
                    (name, email, user["id"]),
                )
            conn.commit()
        except psycopg.IntegrityError:
            conn.close()
            return render_template(
                "edit_profile.html", user=user,
                error="This email is already in use by another account.",
            )
        conn.close()

        return redirect(url_for("dashboard", profile_saved=1))

    return render_template("edit_profile.html", user=user, error=None)

@app.route("/new-session", methods=["POST"])
@login_required
def new_session():
    user = current_user()

    old_info = active_sessions.pop(user["id"], None)
    conn = get_db()
    if old_info:
        conn.execute(
            "UPDATE class_sessions SET ended_at = %s WHERE id = %s",
            (datetime.now().isoformat(), old_info["class_session_id"]),
        )

    settings = get_user_settings(user["id"])
    totp_interval = (
        settings["qr_interval_seconds"]
        if settings["auto_refresh_enabled"]
        else STATIC_QR_INTERVAL_SECONDS
    )

    class_session_id = str(uuid.uuid4())
    secret = pyotp.random_base32()
    active_sessions[user["id"]] = {
        "secret": secret,
        "totp": pyotp.TOTP(secret, interval=totp_interval, digits=8),
        "class_session_id": class_session_id,
        "qr_size": settings["qr_size"],
        "auto_refresh_enabled": settings["auto_refresh_enabled"],
        "refresh_interval": settings["qr_interval_seconds"],
    }

    conn.execute(
        "INSERT INTO class_sessions (id, teacher_id, created_at, ended_at) VALUES (%s, %s, %s, NULL)",
        (class_session_id, user["id"], datetime.now().isoformat()),
    )
    conn.execute("DELETE FROM pending_scans WHERE teacher_id = %s", (user["id"],))
    conn.commit()
    conn.close()

    return redirect(url_for("display"))

@app.route("/end-session", methods=["POST"])
@login_required
def end_session():
    user = current_user()
    info = active_sessions.pop(user["id"], None)

    ended_session_id = ""
    if info:
        ended_session_id = info["class_session_id"]
        conn = get_db()
        conn.execute(
            "UPDATE class_sessions SET ended_at = %s WHERE id = %s",
            (datetime.now().isoformat(), ended_session_id),
        )
        conn.commit()
        conn.close()

    return redirect(url_for("display", ended=ended_session_id))

@app.route("/delete-session/<session_id>", methods=["POST"])
@login_required
def delete_session(session_id):
    user = current_user()

    conn = get_db()
    owner_row = conn.execute(
        "SELECT teacher_id FROM class_sessions WHERE id = %s", (session_id,)
    ).fetchone()

    if owner_row is None or owner_row[0] != user["id"]:
        conn.close()
        return "Bu oturum bulunamadı.", 404

    conn.execute(
        "UPDATE class_sessions SET deleted_at = %s WHERE id = %s",
        (datetime.now().isoformat(), session_id),
    )
    conn.commit()
    conn.close()

    active_info = active_sessions.get(user["id"])
    if active_info and active_info["class_session_id"] == session_id:
        active_sessions.pop(user["id"], None)

    return redirect(url_for("dashboard"))

@app.route("/trash")
@login_required
def trash():
    user = current_user()

    conn = get_db()
    rows = conn.execute(
        "SELECT id, created_at, deleted_at FROM class_sessions "
        "WHERE teacher_id = %s AND deleted_at IS NOT NULL ORDER BY deleted_at DESC",
        (user["id"],),
    ).fetchall()

    trashed_sessions = []
    for sid, created_at, deleted_at in rows:
        count = conn.execute(
            "SELECT COUNT(*) FROM attendance WHERE teacher_id = %s AND class_session_id = %s",
            (user["id"], sid),
        ).fetchone()[0]
        trashed_sessions.append({
            "id": sid,
            "created_at": format_dt(created_at),
            "deleted_at": format_dt(deleted_at),
            "count": count,
        })
    conn.close()

    return render_template("trash.html", user=user, sessions=trashed_sessions)

@app.route("/trash/delete-all", methods=["POST"])
@login_required
def trash_delete_all():
    user = current_user()

    conn = get_db()
    conn.execute(
        "DELETE FROM attendance WHERE teacher_id = %s AND class_session_id IN "
        "(SELECT id FROM class_sessions WHERE teacher_id = %s AND deleted_at IS NOT NULL)",
        (user["id"], user["id"]),
    )
    conn.execute(
        "DELETE FROM pending_scans WHERE teacher_id = %s AND class_session_id IN "
        "(SELECT id FROM class_sessions WHERE teacher_id = %s AND deleted_at IS NOT NULL)",
        (user["id"], user["id"]),
    )
    conn.execute(
        "DELETE FROM class_sessions WHERE teacher_id = %s AND deleted_at IS NOT NULL",
        (user["id"],),
    )
    conn.commit()
    conn.close()

    return redirect(url_for("trash"))


@app.route("/trash/delete/<session_id>", methods=["POST"])
@login_required
def trash_delete_forever(session_id):
    user = current_user()

    conn = get_db()
    owner_row = conn.execute(
        "SELECT teacher_id FROM class_sessions WHERE id = %s AND deleted_at IS NOT NULL",
        (session_id,),
    ).fetchone()

    if owner_row is None or owner_row[0] != user["id"]:
        conn.close()
        return "Bu oturum bulunamadı.", 404

    conn.execute(
        "DELETE FROM attendance WHERE teacher_id = %s AND class_session_id = %s",
        (user["id"], session_id),
    )
    conn.execute(
        "DELETE FROM pending_scans WHERE teacher_id = %s AND class_session_id = %s",
        (user["id"], session_id),
    )
    conn.execute("DELETE FROM class_sessions WHERE id = %s", (session_id,))
    conn.commit()
    conn.close()

    return redirect(url_for("trash"))

@app.route("/display")
@login_required
def display():
    user = current_user()
    info = active_sessions.get(user["id"])
    ended_session_id = request.args.get("ended") or None
    return render_template(
        "display.html",
        active=info is not None,
        ended_session_id=ended_session_id,
        qr_size=info["qr_size"] if info else DEFAULT_QR_SIZE,
        refresh_interval=info["refresh_interval"] if info else QR_INTERVAL_SECONDS,
        auto_refresh_enabled=info["auto_refresh_enabled"] if info else True,
    )

@app.route("/qr.png")
@login_required
def qr_png():
    user = current_user()
    info = active_sessions.get(user["id"])
    if info is None:
        return "Oturum aktif değil.", 403

    token = info["totp"].now()
    attend_url = f"{request.host_url}attend?token={token}&t={user['id']}"

    img = qrcode.make(attend_url)
    raw_buf = io.BytesIO()
    img.save(raw_buf, format="PNG")
    raw_buf.seek(0)

    size = info.get("qr_size", DEFAULT_QR_SIZE)
    resized = Image.open(raw_buf).convert("RGB").resize((size, size), Image.NEAREST)
    buf = io.BytesIO()
    resized.save(buf, format="PNG")
    buf.seek(0)

    return send_file(buf, mimetype="image/png")

@app.route("/attend", methods=["GET"])
def attend_form():
    token = request.args.get("token", "")
    teacher_id_raw = request.args.get("t", "")

    info = None
    if teacher_id_raw.isdigit():
        info = active_sessions.get(int(teacher_id_raw))

    if info is None:
        return render_template(
            "attend.html",
            error="Şu anda aktif bir yoklama oturumu yok. Öğretmeninize danışın.",
            session_id=None,
        )

    if not info["totp"].verify(token, valid_window=1):
        return render_template(
            "attend.html",
            error="QR kodun süresi doldu, lütfen ekrandaki güncel QR kodu tekrar okutun.",
            session_id=None,
        )

    session_id = secrets.token_urlsafe(24)
    conn = get_db()
    cleanup_expired_scans(conn)
    conn.execute(
        "INSERT INTO pending_scans (session_id, teacher_id, class_session_id, created_at) VALUES (%s, %s, %s, %s)",
        (session_id, int(teacher_id_raw), info["class_session_id"], time.time()),
    )
    conn.commit()
    conn.close()

    return render_template("attend.html", error=None, session_id=session_id)

@app.route("/attend", methods=["POST"])
def attend_submit():
    session_id = request.form.get("session_id", "")
    full_name = request.form.get("full_name", "").strip()
    student_id = request.form.get("student_id", "").strip()
    crn = request.form.get("crn", "").strip() or None

    conn = get_db()
    cleanup_expired_scans(conn)

    row = conn.execute(
        "SELECT teacher_id, class_session_id FROM pending_scans WHERE session_id = %s",
        (session_id,),
    ).fetchone()

    if row is None:
        conn.close()
        return render_template(
            "attend.html",
            error="Oturumun süresi doldu ya da geçersiz. Lütfen ekrandaki QR kodu tekrar okutun.",
            session_id=None,
        )

    teacher_id, class_session_id = row

    if not full_name or not student_id:
        conn.close()
        return render_template(
            "attend.html",
            error="Lütfen ad soyad ve öğrenci numarasını eksiksiz girin.",
            session_id=session_id,
        )

    try:
        conn.execute(
            "INSERT INTO attendance (teacher_id, class_session_id, full_name, student_id, timestamp, crn) VALUES (%s, %s, %s, %s, %s, %s)",
            (teacher_id, class_session_id, full_name, student_id, datetime.now().isoformat(), crn),
        )
        conn.commit()
        message = f"Attendance has been recorded for {full_name} ({student_id})."
    except psycopg.IntegrityError:
        message = "This attendance has already been recorded."
    finally:
        conn.execute("DELETE FROM pending_scans WHERE session_id = %s", (session_id,))
        conn.commit()
        conn.close()

    return render_template("success.html", message=message)

@app.route("/report/<session_id>")
@login_required
def report_detail(session_id):
    user = current_user()

    conn = get_db()
    owner_row = conn.execute(
        "SELECT teacher_id, created_at FROM class_sessions WHERE id = %s", (session_id,)
    ).fetchone()

    if owner_row is None or owner_row[0] != user["id"]:
        conn.close()
        return "Bu oturum bulunamadı.", 404

    rows = conn.execute(
        "SELECT full_name, student_id, crn FROM attendance "
        "WHERE teacher_id = %s AND class_session_id = %s ORDER BY timestamp",
        (user["id"], session_id),
    ).fetchall()
    conn.close()

    return render_template(
        "report.html", rows=rows, created_at=format_dt(owner_row[1]), session_id=session_id
    )


@app.route("/report/<session_id>/download")
@login_required
def report_download(session_id):
    user = current_user()

    conn = get_db()
    owner_row = conn.execute(
        "SELECT teacher_id, created_at FROM class_sessions WHERE id = %s", (session_id,)
    ).fetchone()

    if owner_row is None or owner_row[0] != user["id"]:
        conn.close()
        return "Bu oturum bulunamadı.", 404

    rows = conn.execute(
        "SELECT full_name, student_id, crn FROM attendance "
        "WHERE teacher_id = %s AND class_session_id = %s ORDER BY timestamp",
        (user["id"], session_id),
    ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    ws["A1"] = "Session Date:"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = format_dt(owner_row[1])

    headers = ["Full Name", "Student ID", "CRN"]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)

    for full_name, student_id, crn in rows:
        ws.append([full_name, student_id, crn or ""])

    for i, width in enumerate([28, 18, 14], start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"attendance_{session_id}.xlsx",
    )

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)